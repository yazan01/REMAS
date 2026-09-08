"""AI-01 — text and metadata extraction from uploaded supporting documents.

Runs locally with no network call. Each extracted page keeps its number so an
AI finding can cite "document X, page 3" rather than gesturing at a whole file
(AI-03).

Scanned documents (images, image-only PDFs) need OCR. The hook is here and
reports `ocr_required` when a file yields no text layer; wiring an actual OCR
provider is the decision flagged in the implementation plan — its quality on
Arabic documents has to be measured against real client files before a provider
is chosen, so nothing is guessed at here.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

log = logging.getLogger("remas.ai.extraction")

TEXT_MIN_CHARS = 20


def _pdf(path: Path) -> tuple[list[dict], dict]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        pages.append({"page": index, "text": text})
    meta = {k.lstrip("/"): str(v) for k, v in (reader.metadata or {}).items()}
    return pages, {"page_count": len(reader.pages), "document_metadata": meta}


def _docx(path: Path) -> tuple[list[dict], dict]:
    from docx import Document as Docx

    doc = Docx(str(path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return [{"page": 1, "text": "\n".join(paragraphs)}], {"paragraph_count": len(paragraphs)}


def _xlsx(path: Path) -> tuple[list[dict], dict]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    pages = []
    for index, sheet in enumerate(wb.worksheets, start=1):
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(v).strip() for v in row if v not in (None, "")]
            if values:
                rows.append(" | ".join(values))
        pages.append({"page": index, "sheet": sheet.title, "text": "\n".join(rows)})
    wb.close()
    return pages, {"sheet_count": len(pages)}


def _pptx(path: Path) -> tuple[list[dict], dict]:
    from pptx import Presentation

    prs = Presentation(str(path))
    pages = []
    for index, slide in enumerate(prs.slides, start=1):
        chunks = [
            shape.text.strip()
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        ]
        pages.append({"page": index, "text": "\n".join(chunks)})
    return pages, {"slide_count": len(pages)}


EXTRACTORS = {
    "application/pdf": _pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _docx,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": _xlsx,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": _pptx,
}

IMAGE_TYPES = {"image/png", "image/jpeg"}


def extract(path: str | Path, content_type: str) -> dict[str, Any]:
    """Returns {pages, text, char_count, ocr_required, language_hint, meta, error}."""
    file_path = Path(path)
    result: dict[str, Any] = {
        "pages": [],
        "text": "",
        "char_count": 0,
        "ocr_required": False,
        "language_hint": None,
        "meta": {},
        "error": None,
    }

    if not file_path.exists():
        result["error"] = "file_missing"
        return result

    if content_type in IMAGE_TYPES:
        # No text layer exists in an image by definition.
        result["ocr_required"] = True
        result["meta"] = {"kind": "image"}
        return result

    extractor = EXTRACTORS.get(content_type)
    if extractor is None:
        result["error"] = "unsupported_type"
        return result

    try:
        pages, meta = extractor(file_path)
    except Exception as exc:  # noqa: BLE001 - a bad upload must not break the request
        log.warning("extraction failed for %s: %s", file_path.name, exc)
        result["error"] = f"extraction_failed: {type(exc).__name__}"
        return result

    text = "\n\n".join(p["text"] for p in pages if p.get("text"))
    result["pages"] = pages
    result["text"] = text
    result["char_count"] = len(text)
    result["meta"] = meta
    result["language_hint"] = detect_language(text)
    # A PDF that yields almost nothing is a scan, not an empty document.
    result["ocr_required"] = content_type == "application/pdf" and len(text) < TEXT_MIN_CHARS
    return result


def detect_language(text: str) -> str | None:
    """Arabic vs Latin by script share — enough to route prompts (AI-07)."""
    if not text:
        return None
    arabic = sum(1 for ch in text if "؀" <= ch <= "ۿ")
    latin = sum(1 for ch in text if ("a" <= ch.lower() <= "z"))
    total = arabic + latin
    if total < 10:
        return None
    if arabic / total > 0.6:
        return "ar"
    if latin / total > 0.6:
        return "en"
    return "mixed"


def excerpt(pages: list[dict], needle: str, window: int = 240) -> dict | None:
    """Find a phrase and return a citation: page number plus surrounding text."""
    target = needle.strip().lower()
    if not target:
        return None
    for page in pages:
        text = page.get("text") or ""
        position = text.lower().find(target)
        if position >= 0:
            start = max(0, position - window // 2)
            return {
                "page": page.get("page"),
                "excerpt": text[start : start + window].strip(),
            }
    return None
